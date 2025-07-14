import os
from pathlib import Path
from docx import Document

RAW_DIR = Path("data/raw")
VALIDATED_DIR = Path("data/validated")
VALIDATED_DIR.mkdir(parents=True, exist_ok=True)

SEPARATOR = "---"

def get_bold_text_from_cell(cell):
    for paragraph in cell.paragraphs:
        for run in paragraph.runs:
            if run.bold and run.text.strip():
                return run.text.strip()
    return None

def extract_entries_from_tables(doc):
    entries = []
    warnings = []

    for table in doc.tables:
        rows = table.rows
        if not rows or len(rows[0].cells) < 4:
            warnings.append("⚠️ Table is empty or has too few columns — skipped.")
            continue

        first_row = [cell.text.strip().lower() for cell in rows[0].cells]
        has_header = any("chw name" in c or "testimonial" in c or "date" in c for c in first_row)

        if has_header:
            data_rows = rows[1:]
            try:
                idx_gender = first_row.index("chw gender")
                idx_name = first_row.index("chw name")
                idx_date = first_row.index("date")
                idx_testimonial = first_row.index("testimonial")
            except ValueError:
                warnings.append("⚠️ Header format issue — using default indexes (0: gender, 1: name, 2: date, 3: testimonial).")
                idx_gender, idx_name, idx_date, idx_testimonial = 0, 1, 2, 3
        else:
            data_rows = rows
            idx_gender, idx_name, idx_date, idx_testimonial = 0, 1, 2, 3

        for i, row in enumerate(data_rows, 1):
            cells = row.cells
            if len(cells) <= max(idx_testimonial, idx_name, idx_date):
                warnings.append(f"⚠️ Row {i} has insufficient columns — skipped.")
                continue

            gender = cells[idx_gender].text.strip() or "unknown"
            speaker = cells[idx_name].text.strip() or "unknown"
            date = cells[idx_date].text.strip() or "unknown"
            testimonial_cell = cells[idx_testimonial]

            full_text = testimonial_cell.text.strip()
            if not full_text:
                warnings.append(f"⚠️ Missing testimonial in row {i} for '{speaker}' — skipped.")
                continue

            topic = get_bold_text_from_cell(testimonial_cell)
            if not topic:
                topic = full_text.split("\n")[0].strip().split(" ")[0] or "Unknown"
                warnings.append(f"⚠️ No bold topic in row {i} for '{speaker}' — fallback used.")

            body = [p.text.strip() for p in testimonial_cell.paragraphs if p.text.strip()]

            entries.append({
                "topic": topic,
                "speaker": speaker,
                "gender": gender,
                "date": date,
                "body": body
            })

    return entries, warnings

def save_cleaned_doc(entries, output_path):
    doc = Document()
    for entry in entries:
        doc.add_paragraph(f"Topic Title: {entry['topic']}")
        doc.add_paragraph(f"Speaker: {entry['speaker']}")
        doc.add_paragraph(f"Gender: {entry['gender']}")
        doc.add_paragraph(f"Date: {entry['date']}")
        for line in entry["body"]:
            doc.add_paragraph(line)
        doc.add_paragraph(SEPARATOR)
    doc.save(output_path)

def validate_all_docs():
    docx_files = list(RAW_DIR.glob("*.docx"))
    failed = []

    for file_path in docx_files:
        print(f"\n🔍 Validating: {file_path.name}")
        try:
            doc = Document(file_path)
            entries, warnings = extract_entries_from_tables(doc)

            if not entries:
                failed.append(file_path.name)
                print(f"❌ No valid entries found in {file_path.name} — file skipped.\n")
                continue

            output_path = VALIDATED_DIR / f"{file_path.stem}.validated.docx"
            save_cleaned_doc(entries, output_path)

            print(f"✅ {len(entries)} entries saved to {output_path.name}")
            for i, entry in enumerate(entries, 1):
                print(f"  Entry {i}: {entry['topic']} | {entry['speaker']} | {entry['gender']} | {entry['date']}")

            if warnings:
                print("\n⚠️ Warnings:")
                for w in warnings:
                    print(" -", w)
            else:
                print("✅ No structural warnings.")
        except Exception as e:
            failed.append(file_path.name)
            print(f"❌ Error validating {file_path.name}: {e}")

    if failed:
        print("\n📛 The following files failed or had no valid entries:")
        for f in failed:
            print(" -", f)
    else:
        print("\n🎉 All files validated successfully!")

if __name__ == "__main__":
    validate_all_docs()



